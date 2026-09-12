from flask import Flask, render_template, request, redirect, url_for, send_file
import mysql.connector
import os

from config import DB_CONFIG

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet

from io import BytesIO


# =========================================================
# FLASK APP
# =========================================================

app = Flask(__name__)


# =========================================================
# DATABASE CONNECTION
# =========================================================

def get_db_connection():

    connection_config = {
        "host": DB_CONFIG["host"],
        "port": DB_CONFIG["port"],
        "user": DB_CONFIG["user"],
        "password": DB_CONFIG["password"],
        "database": DB_CONFIG["database"],
        "connection_timeout": 10
    }

    # Aiven MySQL SSL connection
    if DB_CONFIG["host"] != "localhost":

        ca_path = os.path.join(
            os.path.dirname(__file__),
            "ca.pem"
        )

        if os.path.exists(ca_path):

            connection_config.update({
                "ssl_ca": ca_path,
                "ssl_verify_cert": True,
                "ssl_verify_identity": True
            })

    return mysql.connector.connect(
        **connection_config
    )


# =========================================================
# LOGIN PAGE
# =========================================================

@app.route("/")
def login_page():

    return render_template("login.html")


# =========================================================
# LOGIN
# =========================================================

@app.route("/login", methods=["POST"])
def login():

    username = request.form["username"]
    password = request.form["password"]

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute(
        """
        SELECT *
        FROM users
        WHERE username = %s
        AND password = %s
        """,
        (username, password)
    )

    user = cursor.fetchone()

    cursor.close()
    conn.close()

    if user:

        return redirect(
            url_for("dashboard")
        )

    return """
    <h2>Invalid Username or Password</h2>
    <a href="/">Back to Login</a>
    """


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
def dashboard():

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    # Total Students
    cursor.execute(
        """
        SELECT COUNT(*) AS total
        FROM students
        """
    )

    total_students = cursor.fetchone()["total"]

    # Present Today
    cursor.execute(
        """
        SELECT COUNT(*) AS present
        FROM attendance
        WHERE attendance_date = CURDATE()
        AND status = 'Present'
        """
    )

    present_today = cursor.fetchone()["present"]

    # Absent Today
    cursor.execute(
        """
        SELECT COUNT(*) AS absent
        FROM attendance
        WHERE attendance_date = CURDATE()
        AND status = 'Absent'
        """
    )

    absent_today = cursor.fetchone()["absent"]

    # Attendance Percentage
    if total_students > 0:

        attendance_percentage = (
            present_today / total_students
        ) * 100

    else:

        attendance_percentage = 0

    cursor.close()
    conn.close()

    return render_template(
        "dashboard.html",
        total_students=total_students,
        present_today=present_today,
        absent_today=absent_today,
        attendance_percentage=round(
            attendance_percentage,
            2
        )
    )


# =========================================================
# STUDENT MANAGEMENT
# =========================================================

@app.route("/students")
def students():

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute(
        """
        SELECT
            id,
            roll_no,
            prn_no,
            name,
            class_name
        FROM students
        ORDER BY CAST(roll_no AS UNSIGNED) ASC
        """
    )

    student_list = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        "students.html",
        students=student_list
    )


# =========================================================
# ADD STUDENT
# =========================================================

@app.route("/add-student", methods=["GET", "POST"])
def add_student():

    if request.method == "POST":

        roll_no = request.form["roll_no"]
        prn_no = request.form["prn_no"]
        name = request.form["name"]
        class_name = request.form["class_name"]

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute(
            """
            INSERT INTO students
            (
                roll_no,
                prn_no,
                name,
                class_name
            )
            VALUES (%s, %s, %s, %s)
            """,
            (
                roll_no,
                prn_no,
                name,
                class_name
            )
        )

        conn.commit()

        cursor.close()
        conn.close()

        return redirect(
            url_for("students")
        )

    return render_template(
        "add_student.html"
    )


# =========================================================
# EDIT STUDENT
# =========================================================

@app.route(
    "/edit-student/<int:student_id>",
    methods=["GET", "POST"]
)
def edit_student(student_id):

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    if request.method == "POST":

        roll_no = request.form.get("roll_no")
        prn_no = request.form.get("prn_no")
        name = request.form.get("name")
        class_name = request.form.get("class_name")

        if not roll_no or not prn_no or not name or not class_name:

            cursor.close()
            conn.close()

            return "All fields are required."

        cursor.execute(
            """
            UPDATE students
            SET
                roll_no = %s,
                prn_no = %s,
                name = %s,
                class_name = %s
            WHERE id = %s
            """,
            (
                roll_no,
                prn_no,
                name,
                class_name,
                student_id
            )
        )

        conn.commit()

        cursor.close()
        conn.close()

        return redirect(
            url_for("students")
        )

    cursor.execute(
        """
        SELECT
            id,
            roll_no,
            prn_no,
            name,
            class_name
        FROM students
        WHERE id = %s
        """,
        (student_id,)
    )

    student = cursor.fetchone()

    cursor.close()
    conn.close()

    if student is None:

        return """
        <h2>Student not found</h2>
        <a href="/students">Back</a>
        """

    return render_template(
        "edit_student.html",
        student=student
    )


# =========================================================
# DELETE STUDENT
# =========================================================

@app.route("/delete-student/<int:id>")
def delete_student(id):

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        DELETE FROM students
        WHERE id = %s
        """,
        (id,)
    )

    conn.commit()

    cursor.close()
    conn.close()

    return redirect(
        url_for("students")
    )


# =========================================================
# MARK ATTENDANCE
# =========================================================

@app.route(
    "/mark-attendance",
    methods=["GET", "POST"]
)
def mark_attendance():

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute(
        """
        SELECT
            id,
            roll_no,
            prn_no,
            name,
            class_name
        FROM students
        ORDER BY CAST(roll_no AS UNSIGNED) ASC
        """
    )

    students = cursor.fetchall()

    # =====================================================
    # SAVE ATTENDANCE
    # =====================================================

    if request.method == "POST":

        attendance_date = request.form[
            "attendance_date"
        ]

        for student in students:

            # Default = Present
            status = request.form.get(
                f"status_{student['id']}",
                "Present"
            )

            cursor.execute(
                """
                INSERT INTO attendance
                (
                    student_id,
                    attendance_date,
                    status
                )
                VALUES (%s, %s, %s)

                ON DUPLICATE KEY UPDATE
                    status = %s
                """,
                (
                    student["id"],
                    attendance_date,
                    status,
                    status
                )
            )

        conn.commit()

        cursor.close()
        conn.close()

        return redirect(
            url_for(
                "view_attendance",
                attendance_date=attendance_date
            )
        )

    cursor.close()
    conn.close()

    return render_template(
        "mark_attendance.html",
        students=students
    )


# =========================================================
# VIEW ATTENDANCE
# =========================================================

@app.route(
    "/view-attendance",
    methods=["GET", "POST"]
)
def view_attendance():

    attendance_date = request.values.get(
        "attendance_date"
    )

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    attendance_list = []

    if attendance_date:

        cursor.execute(
            """
            SELECT
                s.roll_no,
                s.prn_no,
                s.name,
                s.class_name,
                a.status

            FROM students s

            LEFT JOIN attendance a
                ON s.id = a.student_id
                AND a.attendance_date = %s

            ORDER BY
                CAST(
                    s.roll_no AS UNSIGNED
                ) ASC
            """,
            (attendance_date,)
        )

        attendance_list = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        "view_attendance.html",
        attendance=attendance_list,
        selected_date=attendance_date or ""
    )


# =========================================================
# DATE-WISE EXCEL
# =========================================================

@app.route("/download-daily-excel")
def download_daily_excel():

    attendance_date = request.args.get(
        "attendance_date"
    )

    if not attendance_date:

        return "Attendance date is required."

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute(
        """
        SELECT
            students.roll_no,
            students.prn_no,
            students.name,
            students.class_name,
            attendance.status

        FROM students

        LEFT JOIN attendance
            ON students.id = attendance.student_id
            AND attendance.attendance_date = %s

        ORDER BY
            CAST(
                students.roll_no AS UNSIGNED
            ) ASC
        """,
        (attendance_date,)
    )

    attendance_list = cursor.fetchall()

    cursor.close()
    conn.close()

    # =====================================================
    # CREATE EXCEL
    # =====================================================

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Daily Attendance"

    headers = [
        "Date",
        "Roll No",
        "PRN No",
        "Student Name",
        "Class",
        "Status"
    ]

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    row_number = 2

    for row in attendance_list:

        sheet.cell(
            row=row_number,
            column=1
        ).value = attendance_date

        sheet.cell(
            row=row_number,
            column=2
        ).value = row["roll_no"]

        sheet.cell(
            row=row_number,
            column=3
        ).value = row["prn_no"]

        sheet.cell(
            row=row_number,
            column=4
        ).value = row["name"]

        sheet.cell(
            row=row_number,
            column=5
        ).value = row["class_name"]

        sheet.cell(
            row=row_number,
            column=6
        ).value = (
            row["status"]
            if row["status"]
            else "Not Marked"
        )

        row_number += 1

    # Auto Width
    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = max_length + 3

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            f"Attendance_{attendance_date}.xlsx"
        ),
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# =========================================================
# REPORTS PAGE
# =========================================================

@app.route("/reports")
def reports():

    attendance_date = request.args.get(
        "attendance_date"
    )

    attendance_month = request.args.get(
        "attendance_month"
    )

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    report_list = []

    # =====================================================
    # DATE-WISE REPORT
    # =====================================================

    if attendance_date:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,
                attendance.status

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id
                AND attendance.attendance_date = %s

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """,
            (attendance_date,)
        )

        report_list = cursor.fetchall()

    # =====================================================
    # MONTH-WISE REPORT
    # =====================================================

    elif attendance_month:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,

                COUNT(attendance.id)
                    AS total_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Present'
                        THEN 1
                        ELSE 0
                    END
                ) AS present_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Absent'
                        THEN 1
                        ELSE 0
                    END
                ) AS absent_days

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id
                AND DATE_FORMAT(
                    attendance.attendance_date,
                    '%Y-%m'
                ) = %s

            GROUP BY
                students.id,
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """,
            (attendance_month,)
        )

        report_list = cursor.fetchall()

    # =====================================================
    # OVERALL REPORT
    # =====================================================

    else:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,

                COUNT(attendance.id)
                    AS total_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Present'
                        THEN 1
                        ELSE 0
                    END
                ) AS present_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Absent'
                        THEN 1
                        ELSE 0
                    END
                ) AS absent_days

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id

            GROUP BY
                students.id,
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """
        )

        report_list = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template(
        "reports.html",
        reports=report_list,
        selected_date=attendance_date or "",
        selected_month=attendance_month or ""
    )


# =========================================================
# REPORT EXCEL
# DATE-WISE + MONTH-WISE + OVERALL
# =========================================================

@app.route("/download-report-excel")
def download_report_excel():

    attendance_date = request.args.get(
        "attendance_date"
    )

    attendance_month = request.args.get(
        "attendance_month"
    )

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    # =====================================================
    # DATE-WISE EXCEL
    # =====================================================

    if attendance_date:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,
                attendance.status

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id
                AND attendance.attendance_date = %s

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """,
            (attendance_date,)
        )

        report_list = cursor.fetchall()

        headers = [
            "Date",
            "Roll No",
            "PRN No",
            "Student Name",
            "Class",
            "Status"
        ]

    # =====================================================
    # MONTH-WISE EXCEL
    # =====================================================

    elif attendance_month:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,

                COUNT(attendance.id)
                    AS total_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Present'
                        THEN 1
                        ELSE 0
                    END
                ) AS present_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Absent'
                        THEN 1
                        ELSE 0
                    END
                ) AS absent_days

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id
                AND DATE_FORMAT(
                    attendance.attendance_date,
                    '%Y-%m'
                ) = %s

            GROUP BY
                students.id,
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """,
            (attendance_month,)
        )

        report_list = cursor.fetchall()

        headers = [
            "Month",
            "Roll No",
            "PRN No",
            "Student Name",
            "Class",
            "Total Days",
            "Present Days",
            "Absent Days",
            "Attendance %"
        ]

    # =====================================================
    # OVERALL EXCEL
    # =====================================================

    else:

        cursor.execute(
            """
            SELECT
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name,

                COUNT(attendance.id)
                    AS total_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Present'
                        THEN 1
                        ELSE 0
                    END
                ) AS present_days,

                SUM(
                    CASE
                        WHEN attendance.status = 'Absent'
                        THEN 1
                        ELSE 0
                    END
                ) AS absent_days

            FROM students

            LEFT JOIN attendance
                ON students.id = attendance.student_id

            GROUP BY
                students.id,
                students.roll_no,
                students.prn_no,
                students.name,
                students.class_name

            ORDER BY
                CAST(
                    students.roll_no AS UNSIGNED
                ) ASC
            """
        )

        report_list = cursor.fetchall()

        headers = [
            "Roll No",
            "PRN No",
            "Student Name",
            "Class",
            "Total Days",
            "Present Days",
            "Absent Days",
            "Attendance %"
        ]

    cursor.close()
    conn.close()

    # =====================================================
    # CREATE EXCEL
    # =====================================================

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Attendance Report"

    # Headers
    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    row_number = 2

    # =====================================================
    # DATE-WISE DATA
    # =====================================================

    if attendance_date:

        for row in report_list:

            sheet.cell(
                row=row_number,
                column=1
            ).value = attendance_date

            sheet.cell(
                row=row_number,
                column=2
            ).value = row["roll_no"]

            sheet.cell(
                row=row_number,
                column=3
            ).value = row["prn_no"]

            sheet.cell(
                row=row_number,
                column=4
            ).value = row["name"]

            sheet.cell(
                row=row_number,
                column=5
            ).value = row["class_name"]

            sheet.cell(
                row=row_number,
                column=6
            ).value = (
                row["status"]
                if row["status"]
                else "Not Marked"
            )

            row_number += 1

    # =====================================================
    # MONTH-WISE DATA
    # =====================================================

    elif attendance_month:

        for row in report_list:

            total_days = row["total_days"] or 0
            present_days = row["present_days"] or 0
            absent_days = row["absent_days"] or 0

            if total_days > 0:

                percentage = (
                    present_days /
                    total_days
                ) * 100

            else:

                percentage = 0

            sheet.cell(
                row=row_number,
                column=1
            ).value = attendance_month

            sheet.cell(
                row=row_number,
                column=2
            ).value = row["roll_no"]

            sheet.cell(
                row=row_number,
                column=3
            ).value = row["prn_no"]

            sheet.cell(
                row=row_number,
                column=4
            ).value = row["name"]

            sheet.cell(
                row=row_number,
                column=5
            ).value = row["class_name"]

            sheet.cell(
                row=row_number,
                column=6
            ).value = total_days

            sheet.cell(
                row=row_number,
                column=7
            ).value = present_days

            sheet.cell(
                row=row_number,
                column=8
            ).value = absent_days

            sheet.cell(
                row=row_number,
                column=9
            ).value = round(
                percentage,
                2
            )

            row_number += 1

    # =====================================================
    # OVERALL DATA
    # =====================================================

    else:

        for row in report_list:

            total_days = row["total_days"] or 0
            present_days = row["present_days"] or 0
            absent_days = row["absent_days"] or 0

            if total_days > 0:

                percentage = (
                    present_days /
                    total_days
                ) * 100

            else:

                percentage = 0

            sheet.cell(
                row=row_number,
                column=1
            ).value = row["roll_no"]

            sheet.cell(
                row=row_number,
                column=2
            ).value = row["prn_no"]

            sheet.cell(
                row=row_number,
                column=3
            ).value = row["name"]

            sheet.cell(
                row=row_number,
                column=4
            ).value = row["class_name"]

            sheet.cell(
                row=row_number,
                column=5
            ).value = total_days

            sheet.cell(
                row=row_number,
                column=6
            ).value = present_days

            sheet.cell(
                row=row_number,
                column=7
            ).value = absent_days

            sheet.cell(
                row=row_number,
                column=8
            ).value = round(
                percentage,
                2
            )

            row_number += 1

    # =====================================================
    # AUTO COLUMN WIDTH
    # =====================================================

    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = max_length + 3

    # =====================================================
    # SAVE EXCEL
    # =====================================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    # File Name
    if attendance_date:

        filename = (
            f"Attendance_Report_{attendance_date}.xlsx"
        )

    elif attendance_month:

        filename = (
            f"Monthly_Attendance_{attendance_month}.xlsx"
        )

    else:

        filename = (
            "Overall_Attendance_Report.xlsx"
        )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# =========================================================
# REPORT PDF - DATE WISE
# =========================================================

@app.route("/download-report-pdf")
def download_report_pdf():

    attendance_date = request.args.get(
        "attendance_date"
    )

    if not attendance_date:

        return "Attendance date is required."

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    cursor.execute(
        """
        SELECT
            students.roll_no,
            students.prn_no,
            students.name,
            students.class_name,
            attendance.status

        FROM students

        LEFT JOIN attendance
            ON students.id = attendance.student_id
            AND attendance.attendance_date = %s

        ORDER BY
            CAST(
                students.roll_no AS UNSIGNED
            ) ASC
        """,
        (attendance_date,)
    )

    report_list = cursor.fetchall()

    cursor.close()
    conn.close()

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()

    elements = []

    title = Paragraph(
        f"Attendance Report - {attendance_date}",
        styles["Title"]
    )

    elements.append(title)

    elements.append(
        Spacer(1, 15)
    )

    data = [
        [
            "Roll No",
            "PRN No",
            "Student Name",
            "Class",
            "Status"
        ]
    ]

    for row in report_list:

        status = (
            row["status"]
            if row["status"]
            else "Not Marked"
        )

        data.append(
            [
                row["roll_no"],
                row["prn_no"],
                row["name"],
                row["class_name"],
                status
            ]
        )

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.grey
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    1,
                    colors.black
                ),

                (
                    "FONTNAME",
                    (0, 1),
                    (-1, -1),
                    "Helvetica"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                )
            ]
        )
    )

    elements.append(table)

    document.build(elements)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            f"Attendance_Report_{attendance_date}.pdf"
        ),
        mimetype="application/pdf"
    )


# =========================================================
# TEST DATABASE
# =========================================================

@app.route("/test-db")
def test_db():

    try:

        conn = get_db_connection()

        cursor = conn.cursor()

        cursor.execute(
            "SELECT 1"
        )

        result = cursor.fetchone()

        cursor.close()
        conn.close()

        return f"""
        <h2>Database Connected Successfully!</h2>
        <p>Result: {result}</p>
        """

    except Exception as e:

        return f"""
        <h2>Database Connection Failed</h2>
        <p>{e}</p>
        """


# =========================================================
# RUN APPLICATION
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )